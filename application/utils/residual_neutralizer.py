"""v62.13.9.4.20 — NEUTRALISATION DE DERNIER RECOURS (décision Robin du 04/09/2026).

Depuis la 9.4.18, chaque run réel a été retenu sur UNE ligne, chaque fois
d'une forme nouvelle (« Stabilité 1 run », un remplacement chiffré du registre,
une unité tronquée, l'horodatage, « 3 claim(s) », « 70/100 », « 10/2026 », une
alerte méthodologique citant un chiffre pour le récuser). Le contrat des
fixtures ferme chaque cas pour toujours — après. Ce module ferme la queue
de distribution : quand le contrôle final trouve au plus RESIDUAL_TOLERANCE
unités résiduelles, elles sont NEUTRALISÉES dans le fichier (réécriture
qualitative sans chiffre, sinon marque de retrait), l'original entre au
registre embarqué sous tag, et le contrôle est rejoué une fois. Propre →
livré, avec un avertissement visible et tracé. Encore fautif, ou plus de
RESIDUAL_TOLERANCE → blocage, comme avant. Rien de non étayé ne sort ; l'export
n'est plus otage d'une ligne.
"""
from __future__ import annotations

import io
import re
from typing import Any, Iterable

# v62.13.9.4.20.2 — PLUS DE PLAFOND (décision Robin du 05/09/2026). Chaque
# résidu est neutralisé dans le fichier et archivé au registre : rien de non
# étayé ne sort, quel que soit leur nombre. Le plafond de trois était une
# prudence, pas une doctrine, et c'est lui qui laissait l'utilisateur sans
# document. Au-delà de RESIDUAL_WARNING résidus, le dossier porte un
# avertissement — le run est à revoir, le document sort quand même.
RESIDUAL_TOLERANCE = None
RESIDUAL_WARNING = 3

_NOTE_FR = ("Note de sortie : des lignes chiffrées non étayées ont été neutralisées au dernier contrôle "
            "et conservées dans le registre embarqué en annexe (texte original, statut, décision requise).")
_NOTE_EN = ("Exit note: unsupported numbered lines were neutralised at the final check and kept in the "
            "embedded register in the appendix (original text, status, decision required).")


def _norm(s: str) -> str:
    # v62.13.9.4.31.1 — même normalisation que l'oracle (NFC, espaces insécables,
    # sauts de ligne → un espace) : contrôle et réparation lisent le même texte.
    import unicodedata as _ud
    t = _ud.normalize("NFC", str(s or "")).replace("\u00A0", " ").replace("\u202F", " ")
    return re.sub(r"\s+", " ", t).strip()


def _neutral_text(text: str, lang: str) -> str:
    """Réécriture sans chiffre ; à défaut, marque de retrait."""
    try:
        from utils.editorial_quantity import rewrite_quantities
        out = rewrite_quantities(text, lang)
        if out and not re.search(r"\d", out) and _norm(out) != _norm(text):
            return out
    except Exception:
        pass
    return ("(phrase chiffrée retirée du corps — conservée au registre du dossier complet)"
            if not str(lang).lower().startswith("en") else
            "(numbered sentence withdrawn from the body — kept in the register of the complete file)")


def _all_paragraph_texts(doc) -> list[str]:
    out = [_norm(p.text) for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                out.extend(_norm(p.text) for p in cell.paragraphs)
    return out


def _set_paragraph_text(p, text: str) -> None:
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = text


_SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?;])\s+(?=[A-ZÀ-ÝÉÈÊÎÔÛ«(\d])")


def _mask_non_magnitudes(text: str) -> str:
    """Dates, repères de période, numéros d'ordre et comptes de structure ne
    sont pas des grandeurs — même lecture que le typeur (v62.13.9.4.32)."""
    try:
        from utils.declared_quantities import (_CALENDAR_DATE_RE, _PERIOD_LABEL_RE,
                                               _ORDINAL_LABEL_RE, _STRUCTURE_COUNT_RE)
        t = _CALENDAR_DATE_RE.sub(lambda m: " " * len(m.group(0)), text)
        t = _PERIOD_LABEL_RE.sub(lambda m: " " * len(m.group(0)), t)
        t = _ORDINAL_LABEL_RE.sub(lambda m: m.group(1) + " " * len(m.group(2)), t)
        t = _STRUCTURE_COUNT_RE.sub(lambda m: " " * len(m.group(0)), t)
        t = re.sub(r"\b(?:19|20)\d{2}\b", "    ", t)
        return t
    except Exception:
        return text


def split_sentences(text: str) -> list[str]:
    return [x for x in _SENTENCE_SPLIT_RE.split(str(text or "").strip()) if x]


def sentence_level_neutral(text: str, lang: str, quantities: Iterable[str] = ()) -> tuple[str, list[str]] | None:
    """v62.13.9.4.32 (A) — LA PHRASE QUI PORTE LE CHIFFRE, pas le paragraphe.

    Run réel ee08f0 (11/09/2026) : un faux positif sur « (T1 2027) » a fait
    remplacer par une note de retrait TOUTE la phrase de recommandation — cinq
    lignes, dont l'orientation elle-même — dans les deux Word ; la trame est
    retombée sur la première option. Un contrôle qui retient une grandeur ne
    doit retirer que ce qui la porte (contre-audit du dossier industriel :
    « ne neutraliser que le fragment réellement contestable »). Désigne les
    phrases portant une grandeur signalée (``quantities``), à défaut celles
    portant un chiffre qui n'est ni date, ni repère, ni numéro d'ordre ; les
    autres phrases restent intactes. Renvoie None quand la découpe n'a pas de
    sens (une seule phrase, ou toutes concernées) : l'appelant garde alors le
    comportement au paragraphe.
    """
    sents = split_sentences(text)
    if len(sents) <= 1:
        return None
    qs = [str(q).strip() for q in (quantities or ()) if str(q or "").strip()]

    def _carries(sent: str) -> bool:
        for q in qs:
            if re.search(r"(?<![\w.,])" + re.escape(q) + r"(?![\w])", _mask_non_magnitudes(sent)):
                return True
        return False
    idx = [i for i, sn in enumerate(sents) if _carries(sn)]
    if not idx:
        idx = [i for i, sn in enumerate(sents) if re.search(r"\d", _mask_non_magnitudes(sn))]
    if not idx or len(idx) == len(sents):
        return None
    out = [(_neutral_text(sn, lang) if i in idx else sn) for i, sn in enumerate(sents)]
    return " ".join(out), [sents[i] for i in idx]


def neutralize_docx_residuals(payload: bytes, units: Iterable[str], lang: str = "fr",
                              quantities: dict | None = None) -> tuple[bytes, list[dict]]:
    """Neutralise chaque unité résiduelle (paragraphe ou cellule dont le texte
    normalisé contient l'unité) et l'archive au registre embarqué.

    v62.13.9.4.32 (A) — ``quantities`` : grandeur(s) signalée(s) par unité
    normalisée ; la neutralisation vise alors la phrase qui la porte."""
    from docx import Document
    try:
        from utils.quantity_ledger import SCN_REGISTER_TAG_FR, SCN_REGISTER_TAG_EN
    except Exception:
        SCN_REGISTER_TAG_FR, SCN_REGISTER_TAG_EN = "[SIGNALÉ — en place, à vérifier]", "[FLAGGED — in place, to verify]"
    en = str(lang).lower().startswith("en")
    tag = SCN_REGISTER_TAG_EN if en else SCN_REGISTER_TAG_FR
    doc = Document(io.BytesIO(payload))
    targets = [_norm(u) for u in units if _norm(u)]
    done: list[dict] = []

    # v62.13.9.4.31.1 — LE PASSAGE EXACT, jamais un texte ressemblant.
    # Jusqu'ici la correspondance se faisait par sous-chaîne (les 120 premiers
    # caractères de l'unité dans le paragraphe, ou le paragraphe dans l'unité) :
    # au run réel du 11/09/2026, l'unité fautive transmise tronquée à 160
    # caractères a fait remplacer une cellule d'inventaire (le nom d'une pièce
    # jointe, contenue dans l'aperçu) qui ne portait pas le chiffre, et le
    # paragraphe réellement signalé restait intact — le contrôle retrouvait la
    # même alerte et bloquait. Le contrôle transporte désormais l'unité entière
    # (``unit_full``) et la désignation suit trois paliers, tous EXACTS :
    #   1. le paragraphe EST l'unité (texte normalisé égal) ; dès qu'un tel
    #      paragraphe existe dans le fichier, rien d'autre n'est touché ;
    #   2. sinon, le paragraphe CONTIENT l'unité entière (reprise verbatim
    #      d'une ligne retenue au sein d'un paragraphe plus long) ;
    #   3. sinon, l'unité est le paragraphe plus sa décoration (contre-audit
    #      v22 §5, 9.4.23 : « • 2. Optimisation … » signalé, pastille « 2 » et
    #      texte dans deux formes) — le paragraphe doit faire au moins la
    #      moitié de l'unité et douze caractères : une cellule courte contenue
    #      dans une longue unité n'est jamais le passage.
    _all_norm = set(_all_paragraph_texts(doc))
    _exact_targets = {u for u in targets if u in _all_norm}

    def _match(text: str) -> str | None:
        t = _norm(text)
        if len(t) < 4 or not re.search(r"[A-Za-zÀ-ÿ]", t):
            return None                                # v62.13.9.4.20.1 — jamais un paragraphe vide ou sans mot
        for u in targets:
            if u == t:
                return u
            if u in _exact_targets:
                continue                               # le passage exact existe ailleurs : rien d'autre
            if len(u) >= 12 and u in t:
                return u
            if len(t) >= 12 and t in u and 2 * len(t) >= len(u):
                return u
        return None

    _qmap = {_norm(k): list(v if isinstance(v, (list, tuple, set)) else [v])
             for k, v in (quantities or {}).items() if _norm(k)}

    def _visit(p, where: str) -> None:
        u = _match(p.text)
        if u is None:
            return
        original = p.text
        granular = sentence_level_neutral(original, lang, _qmap.get(u) or ())
        if granular is not None:
            new_text, withdrawn = granular
            _set_paragraph_text(p, new_text)
            done.append({"where": where, "granularity": "sentence",
                         "original": " ".join(withdrawn)[:400], "paragraph": original[:400],
                         "after": p.text[:200]})
        else:
            _set_paragraph_text(p, _neutral_text(original, lang))
            done.append({"where": where, "granularity": "paragraph",
                         "original": original[:400], "after": p.text[:200]})

    try:
        from utils.system_marks import is_system_paragraph as _isp
    except Exception:
        _isp = lambda _p: False   # noqa: E731
    for p in doc.paragraphs:
        if any(_k in p.text for _k in ("hors corps]", "off-body]", "en place, à vérifier]", "in place, to verify]")) or _isp(p):
            continue                                   # jamais le registre ni l'outil lui-même
        # v62.13.9.4.31 — un titre de rubrique de PREMIER NIVEAU (celui que porte le
        # sommaire et que visent les renvois) n'est jamais remplacé par un message
        # de contrôle. Les sous-titres écrits par les agents restent neutralisables :
        # sans quoi le filet déclarerait des unités « introuvables » et l'export
        # tomberait alors que la ligne existe bien.
        try:
            _st31 = str(getattr(p.style, "name", "") or "").lower()
            if _st31.startswith(("heading 1", "heading 2", "titre 1", "titre 2")) or _st31 in ("title", "titre"):
                continue
        except Exception:
            pass
        _visit(p, "docx:p")
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if not _isp(p):
                        _visit(p, "docx:cell")
    if done:
        # registre embarqué : original sous tag + décision requise
        doc.add_paragraph(("Complément v62.13.8 / v62.13.9.4.20 — Registre embarqué : lignes neutralisées au dernier contrôle"
                           if not en else
                           "Complément v62.13.8 / v62.13.9.4.20 — Embedded register: lines neutralised at the final check"))
        for i, d in enumerate(done, 1):
            doc.add_paragraph(f"SCN-V62139-X{i:03d} · {tag} {_norm(d['original'])}")
            doc.add_paragraph(("Statut de preuve : non validé — retenu par le contrôle final de sortie"
                               if not en else "Proof status: unvalidated — withheld by the final exit check"))
            doc.add_paragraph(("Source : " if not en else "Source: ") + d["where"])
            doc.add_paragraph(("Décision requise : confirmer la baseline, la formule, la provenance et le porteur ; "
                               "réintégrer ou retirer." if not en else
                               "Decision required: confirm baseline, formula, provenance and owner; reinstate or drop."))
        # note visible en tête (sans chiffre)
        first = doc.paragraphs[0] if doc.paragraphs else None
        note = _NOTE_EN if en else _NOTE_FR
        if first is not None:
            new_p = first.insert_paragraph_before(note)
            try:
                new_p.runs[0].italic = True
            except Exception:
                pass
        else:
            doc.add_paragraph(note)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), done


def neutralize_pptx_residuals(payload: bytes, units: Iterable[str], lang: str = "fr",
                              parts: Iterable[str] | None = None) -> tuple[bytes, list[dict]]:
    """Même règle sur le deck : le paragraphe fautif est réécrit sans chiffre.

    v62.13.9.4.23 — contre-audit v22 (§5) : l'appariement par sous-chaîne
    remplaçait une pastille « 2 » par une phrase entière. Un paragraphe
    n'est neutralisé que s'il est ÉGAL à l'unité signalée ou la contient, s'il
    porte au moins quatre caractères et un mot, et — quand le contrôle a nommé
    la slide (« ppt/slides/slide4.xml ») — seulement sur cette slide.
    v62.13.9.4.31.1 — mêmes trois paliers exacts que le DOCX : égal ; sinon
    contient l'unité entière ; sinon est l'unité moins sa décoration (au moins
    la moitié de l'unité). Une forme courte contenue dans une longue unité
    n'est jamais le passage."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(payload))
    targets = [_norm(u) for u in units if _norm(u)]
    slide_ids: set[int] = set()
    for pt in (parts or ()):
        m = re.search(r"slide(\d+)\.xml", str(pt or ""))
        if m:
            slide_ids.add(int(m.group(1)))
    done: list[dict] = []
    _all_norm_pptx: set[str] = set()
    for _slide in prs.slides:
        for _shape in _slide.shapes:
            if getattr(_shape, "has_text_frame", False):
                _all_norm_pptx.update(_norm(_p.text) for _p in _shape.text_frame.paragraphs)
            if getattr(_shape, "has_table", False):
                for _row in _shape.table.rows:
                    for _cell in _row.cells:
                        _all_norm_pptx.update(_norm(_p.text) for _p in _cell.text_frame.paragraphs)
    _exact_pptx = {u for u in targets if u in _all_norm_pptx}

    def _match_pptx(t: str) -> bool:                    # v62.13.9.4.31.1 — trois paliers exacts (cf. DOCX)
        for u in targets:
            if u == t:
                return True
            if u in _exact_pptx:
                continue
            if len(u) >= 12 and u in t:
                return True
            if len(t) >= 12 and t in u and 2 * len(t) >= len(u):
                return True
        return False

    for si, slide in enumerate(prs.slides, 1):
        if slide_ids and si not in slide_ids:
            continue
        for shape in slide.shapes:
            frames = []
            if getattr(shape, "has_text_frame", False):
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        frames.append(cell.text_frame)
            for tf in frames:
                for p in tf.paragraphs:
                    t = _norm(p.text)
                    if len(t) < 4 or not re.search(r"[A-Za-zÀ-ÿ]", t):
                        continue                          # jamais une pastille ou un numéro
                    if _match_pptx(t):
                        original = p.text
                        _set_paragraph_text(p, _neutral_text(original, lang))
                        done.append({"where": f"pptx:slide{si}", "original": original[:400], "after": p.text[:200]})
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), done
