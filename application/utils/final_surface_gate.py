"""final_surface_gate — v62.13.9.4.6

Dernier contrôle indépendant : relit les OCTETS des DOCX/PPTX finaux.

Deux moteurs de conception différente (contre-audit 29/08/2026, P1-05) :
  L — quantity_ledger : deny-by-default typé. Toute quantité sans type
      descriptif certain (whitelist fermée) retient la ligne. C'est le socle.
  A — editorial_quantity : signaux lexicaux (verbes, morphologie, cibles).
      A AFFINE : il ne peut plus inculper une ligne dont toutes les
      quantités portent un type certain (le typage EST la preuve
      descriptive) ; il conserve son autorité sur les lignes sans
      quantité typée (placeholders, cibles qualitatives).

Blocage = L ∪ A(restreint). Un claim strictement lié (claim_binding) est
intégré au typage (CLAIM_BOUND) ; le filet _binding_supports ne s'applique
qu'aux résidus A-seuls.
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Mapping
from xml.etree import ElementTree as ET

from utils import editorial_quantity as _EQ
from utils import quantity_ledger as _QL
from utils.claim_binding import supports as _binding_supports

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _claims(result: Mapping[str, Any] | None) -> list:
    """Preuves numériques du run — v62.13.9.4.15 : les FAITS DE REGISTRE
    (claims REGISTRY-*) en font partie ici aussi. Le gate de surface était
    le dernier juge à ne pas les voir : « structure solo (0 salarié) »,
    libéré par la réparation et par l'oracle, restait bloqué ici avec
    ``untyped_quantity:0`` (rejeu du run réel du 02/09/2026 15:04). Tous
    les juges voient les mêmes preuves, sinon l'un bloque ce que les
    autres ont libéré."""
    base = _claims_base(result)
    try:
        from utils.quantity_ledger import registry_claims as _rc62415
        extra = list(_rc62415(result) or [])
    except Exception:
        extra = []
    return base + extra


def _claims_base(result: Mapping[str, Any] | None) -> list:
    if not isinstance(result, Mapping):
        return []
    return list(result.get("numeric_support_claims") or result.get("canonical_claims") or [])


def _reference_year(result: Mapping[str, Any] | None) -> int | None:
    try:
        from utils.v62136_fr_path_maintenance import _v6213944_reference_year
        return _v6213944_reference_year(result)
    except Exception:
        return None


def _authority(result: Mapping[str, Any] | None) -> list:
    """v62.13.9.4.7 — items d'autorité structurés {value, unit, semantic_role}."""
    return _QL.build_authority_items(result)


def _authority_legacy_unused(result: Mapping[str, Any] | None) -> set[str]:
    try:
        from utils.v62136_fr_path_maintenance import _v6213945_authoritative_values
        return _v6213945_authoritative_values(result)
    except Exception:
        return set()


_QUALIFIED_RE = re.compile(
    r"—\s*[a-zà-ÿ' ]{2,40}à\s+(?:valider|dimensionner|arbitrer|définir|definir)\b"
    r"|\bà\s+(?:dimensionner|arbitrer)\b|\bDimensionnement à valider\b"
    r"|—\s*(?:estimate|estimation)?,?\s*to (?:verify|validate|size)\b"
    r"|\bestimate, to verify\b|\bto be (?:validated|sized|verified)\b",
    re.I,
)

_FR_HINT_RE = re.compile(r"[àâéèêëîïôûùüç]|\b(?:le|la|les|des|de|du|une?|et|dans|pour|sur|avec)\b", re.I)


def _line_lang(line: str, result: Mapping[str, Any] | None) -> str:
    base = str((result or {}).get("language") or "").lower()
    if _FR_HINT_RE.search(line):
        return "fr"
    if re.search(r"\b(?:the|of|and|to|in|with|for)\b", line, re.I):
        return "en"
    return base if base in ("fr", "en") else "fr"


def detector_reasons(line: str, result: Mapping[str, Any] | None = None) -> list[str]:
    """Moteurs L (socle typé) et A (affinage lexical) — tout signal bloque."""
    # v62.13.9.4.5 — idempotence : une ligne portant la qualification produit
    # et sans plus aucun chiffre est une sortie du système, pas une cible.
    if _QUALIFIED_RE.search(line) and not re.search(r"\d", line):
        return []
    if "\u2063" in line:
        return []                            # v62.13.9.4.21 — marqueur d'outil structurel
    try:
        from utils.quantity_ledger import is_status_tagged as _ist
        if _ist(line):
            return []                        # v62.13.9.4.21 — statut en marge
    except Exception:
        pass
    # v62.13.9.4.15.2 — bandeau d'horizon du brief : même autorité que la
    # maintenance et l'oracle (parité des trois juges).
    try:
        from utils.final_content_oracle import _is_brief_horizon_banner as _ibhb
        if _ibhb(line, result):
            return []
    except Exception:
        pass
    try:                                   # v62.13.9.4.17 — registre embarqué tagué ; 9.4.18 — benchmark étiqueté
        from utils.quantity_ledger import (is_tagged_register_line as _itrl, is_labelled_benchmark_line as _ilbl,
                                           is_epistemic_caveat_line as _iecl)
        if _itrl(line) or _ilbl(line) or _iecl(line):
            return []
    except Exception:
        pass
    authority = _authority(result)
    claims = _claims(result) or None
    ctx = {"current_year": _reference_year(result)} if _reference_year(result) else None
    verdict = _QL.classify_line(
        line, _line_lang(line, result), claims=claims,
        authority_values=authority, context=ctx)
    if not verdict.withhold and any(
            r in ("evidence_row", "scenario_register_row", "financial_model_row")
            for r in verdict.reasons):
        return []  # couche de preuve/registre/modèle : hors surface exécutive
    engines: list[str] = []
    if verdict.withhold:
        engines.append("L")
    fully_typed = bool(verdict.typed) and not verdict.untyped
    if not fully_typed and _EQ.unsupported_quantity(
            line, reference_year=_reference_year(result), authoritative_values=authority):
        engines.append("A")
    return engines


def _scan(lines: list[tuple[str, str]], result: Mapping[str, Any] | None) -> list[dict[str, str]]:
    claims = _claims(result)
    residuals: list[dict[str, str]] = []
    for where, text in lines:
        line = re.sub(r"\s+", " ", str(text or "")).strip()
        if not line:
            continue
        engines = detector_reasons(line, result)
        if not engines:
            continue
        if engines == ["A"] and _binding_supports(line, claims):
            continue
        ctx = {"current_year": _reference_year(result)} if _reference_year(result) else None
        verdict = _QL.classify_line(line, _line_lang(line, result), claims=claims or None,
                                    authority_values=_authority(result), context=ctx)
        reason = ";".join(verdict.reasons) if verdict.reasons else (
            _EQ.unsupported_quantity(line, reference_year=_reference_year(result),
                                     authoritative_values=_authority(result)) or "lexical_signal")
        residuals.append({"where": where, "line": line[:200], "line_full": line, "reason": reason,
                          "engines": "+".join(engines)})
    return residuals


_ANNEX_TITLE_RE = re.compile(r"^\s*(?:annexes?|appendix|appendices)\b", re.I)
# Titres canoniques d'annexe émis par le produit lui-même : leur texte exact
# en tête de paragraphe suffit, quel que soit le style (v62.13.9.4.6).
_PRODUCT_ANNEX_RE = re.compile(
    r"^\s*(?:Unvalidated quantified working scenarios|"
    r"Sc[ée]narios chiffr[ée]s de travail non valid[ée]s)\b", re.I)
_HEADING_STYLE_RE = re.compile(r"Heading|Titre|Title", re.I)


def _docx_lines(data: bytes) -> list[tuple[str, str]]:
    """Corps avant annexe + tableaux. v62.13.9.4.6 (P2-03) : la frontière
    d'annexe exige un PARAGRAPHE DE TITRE (style Heading/Titre) commençant
    par « Annexe(s) »/« Appendix » — une mention narrative du mot ne coupe
    plus le scan."""
    lines: list[tuple[str, str]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    body = root.find(f"{_W}body")
    if body is None:
        return lines
    annex = False
    table_paragraphs = {id(p) for tbl in body.iter(f"{_W}tbl") for p in tbl.iter(f"{_W}p")}
    # v62.13.9.4.49 — tableaux écrits par le produit (registre des actions, chronologie multi-séances) : leurs rangées
    # ne sont pas du texte d'agent (audit du 29/09/2026 : une rangée du registre signalée en résidu)
    _product_rows49 = set()
    _hdr_sets49 = ({"topic", "action", "owner", "deadline"}, {"sujet", "action", "porteur", "échéance"},
                   {"date", "new", "carried"}, {"date", "nouvelles", "reprises"}, {"date", "nouvelles", "reconduites"})
    for _tbl49 in body.iter(f"{_W}tbl"):
        _trs49 = list(_tbl49.iter(f"{_W}tr"))
        if not _trs49:
            continue
        _names49 = {re.sub(r"[^\wéèêàç ]", "", "".join(n.text or "" for n in c.iter(f"{_W}t")).replace("\u2063", "")).strip().lower()
                    for c in _trs49[0].iter(f"{_W}tc")}
        if any(h <= _names49 for h in _hdr_sets49):
            _product_rows49.update(id(tr) for tr in _trs49)
    # v62.13.9.4.21 — paragraphes typés « CoDIR System » (contenu de l'outil) : jamais lus.
    def _is_sys(el) -> bool:
        try:
            ps = el.find(f"{_W}pPr/{_W}pStyle")
            return ps is not None and str(ps.get(f"{_W}val") or "") in ("CoDIRSystem", "CoDIR System")
        except Exception:
            return False
    for element in body.iter():
        if element.tag == f"{_W}p":
            if id(element) in table_paragraphs:
                continue  # v62.13.9.4.6 — couvert par la ligne docx:row jointe
            if _is_sys(element):
                continue  # v62.13.9.4.21 — contenu de l'outil, typé à la construction
            text = "".join(node.text or "" for node in element.iter(f"{_W}t"))
            clean = re.sub(r"\s+", " ", text).strip()
            if not clean:
                continue
            style = ""
            ppr = element.find(f"{_W}pPr")
            if ppr is not None:
                pstyle = ppr.find(f"{_W}pStyle")
                if pstyle is not None:
                    style = pstyle.get(f"{_W}val", "")
            if _PRODUCT_ANNEX_RE.match(clean) or (
                    _ANNEX_TITLE_RE.match(clean)
                    and (_HEADING_STYLE_RE.search(style) or clean.isupper())):
                annex = True
            if not annex:
                lines.append(("docx:p", clean))
        elif element.tag == f"{_W}tr" and not annex:
            if id(element) in _product_rows49:
                continue  # v62.13.9.4.49 — rangée d'un tableau écrit par le produit
            if any(_is_sys(pp) for pp in element.iter(f"{_W}p")):
                continue  # v62.13.9.4.21 — rangée d'un tableau de l'outil (en-tête version/date, faits extraits)
            cells = ["".join(node.text or "" for node in cell.iter(f"{_W}t")).strip() for cell in element.iter(f"{_W}tc")]
            row = " | ".join(c for c in cells if c)
            if row:
                lines.append(("docx:row", row))
    return lines


def _pptx_lines(data: bytes) -> list[tuple[str, str]]:
    lines: list[tuple[str, str]] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = sorted(n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        for name in names:
            root = ET.fromstring(archive.read(name))
            short = name.rsplit("/", 1)[-1]
            table_paragraphs = {id(p) for tbl in root.iter(f"{_A}tbl") for p in tbl.iter(f"{_A}p")}
            for row in root.iter(f"{_A}tr"):
                cells = []
                for cell in row.iter(f"{_A}tc"):
                    txt = " ".join("".join(node.text or "" for node in par.iter(f"{_A}t"))
                                   for par in cell.iter(f"{_A}p")).strip()
                    if txt:
                        cells.append(re.sub(r"\s+", " ", txt))
                if cells:
                    lines.append((f"{short}:row", " | ".join(cells)))
            for paragraph in root.iter(f"{_A}p"):
                if id(paragraph) in table_paragraphs:
                    continue  # v62.13.9.4.6 — couvert par la ligne :row jointe
                text = "".join(node.text or "" for node in paragraph.iter(f"{_A}t"))
                clean = re.sub(r"\s+", " ", text).strip()
                if clean:
                    lines.append((short, clean))
    return lines


def audit_docx_bytes(data: bytes, *, result: Mapping[str, Any] | None = None) -> list[dict[str, str]]:
    """Relit les octets du DOCX final (corps avant annexe + tableaux) — L∪A."""
    try:
        return _scan(_docx_lines(data), result)
    except Exception as exc:  # fail-closed : un gate illisible est un résidu
        return [{"where": "docx", "line": f"gate_error: {type(exc).__name__}: {exc}", "reason": "gate_error", "engines": "L+A"}]


def audit_pptx_bytes(data: bytes, *, result: Mapping[str, Any] | None = None) -> list[dict[str, str]]:
    """Relit les octets du PPTX final (cadres + tables des slides) — L∪A."""
    try:
        return _scan(_pptx_lines(data), result)
    except Exception as exc:
        return [{"where": "pptx", "line": f"gate_error: {type(exc).__name__}: {exc}", "reason": "gate_error", "engines": "L+A"}]



# ── v62.13.9.4.17 — AUDIT DE QUALITÉ FINAL (contre-audit ChatGPT du 03/09/2026,
# point 9). L'oracle ne juge que les résidus quantitatifs : il restait vert
# devant une contradiction d'effectif, des fragments cassés et une densité de
# placeholders qui rendait le document illisible. Cet audit dit ce que
# l'oracle ne dit pas — et ne bloque que sur l'incohérence d'entité et
# l'absence du registre embarqué quand des lignes ont été retenues.

_QA_PLACEHOLDER_RE = re.compile(
    r"volume à valider|à dimensionner|à arbitrer|calendrier à arbitrer|\[NON SOURCÉ\]|"
    r"[ée]l[ée]ment chiffr[ée] retir[ée]|phrase chiffr[ée]e retir[ée]e|"
    r"volume to validate|to size|to arbitrate|\[NOT SOURCED\]", re.I)
_QA_HEADCOUNT_BAND_RE = re.compile(
    r"(?:tranche registre|registry band)\s*:\s*(\d[\d\s\u00a0]*)\s*[–\-]\s*(\d[\d\s\u00a0]*)\s*(?:salari[ée]s?|employees?)", re.I)


def quality_audit_docx(payload: bytes, result: Mapping[str, Any] | None) -> dict:
    """Retourne {blocking: [...], warnings: [...], metrics: {...}}."""
    from docx import Document
    import io
    try:
        from utils.deck_layout_guard import repair_dangling_fragments_text as _rdf
    except Exception:
        _rdf = None
    doc = Document(io.BytesIO(payload))
    try:
        from utils.system_marks import is_system_paragraph as _isp
    except Exception:
        _isp = lambda _p: False   # noqa: E731
    paras = [p.text for p in doc.paragraphs if (p.text or "").strip() and not _isp(p)]
    for tbl in doc.tables:
        for row in tbl.rows:
            for c in row.cells:
                if (c.text or "").strip():
                    paras.append(c.text)
    n = max(1, len(paras))
    joined = "\n".join(paras)
    from utils.quantity_ledger import REGISTER_HEADINGS_RE as _RHR49
    _mh49 = _RHR49.search(joined)
    body = joined[:_mh49.start()] if _mh49 else joined          # v62.13.9.4.49 — anciens et nouveaux titres
    blocking: list[str] = []
    warnings: list[str] = []
    metrics: dict[str, Any] = {"paragraphs": n}

    # 1. cohérence d'entité / effectif
    prof = (result or {}).get("company_profile") or {}
    hc = prof.get("headcount") if isinstance(prof, Mapping) else None
    bad_bands = []
    for m in _QA_HEADCOUNT_BAND_RE.finditer(joined):
        try:
            lo, hi = int(re.sub(r"\D", "", m.group(1))), int(re.sub(r"\D", "", m.group(2)))
        except ValueError:
            continue
        if isinstance(hc, (int, float)) and not (lo <= int(hc) <= hi):
            bad_bands.append(m.group(0)[:80])
    metrics["headcount_contradictions"] = len(bad_bands)
    if bad_bands:
        blocking.append(f"effectif contradictoire avec le profil ({hc} salarié(s)) : " + " | ".join(bad_bands[:3]))

    # 2. placeholders : densité
    ph = len(_QA_PLACEHOLDER_RE.findall(body))
    metrics["placeholders"] = ph
    metrics["placeholder_density_per_100"] = round(100.0 * ph / n, 1)
    if ph / n > 0.25:
        warnings.append(f"densité de placeholders élevée : {ph} sur {n} paragraphes ({100.0 * ph / n:.0f} %)")

    # 3. fragments grammaticaux (mêmes règles que le deck)
    frags = []
    if _rdf is not None:
        _tail = lambda x: re.sub(r"[\s.,;:–—\-|]+$", "", x)   # noqa: E731
        for ptxt in paras:
            s = re.sub(r"\s+", " ", ptxt).strip()
            if s and len(s) < 220 and re.search(r"[A-Za-zÀ-ÿ]", s) and _tail(_rdf(s)) != _tail(s):
                frags.append(s[:100])
    metrics["dangling_fragments"] = len(frags)
    if frags:
        warnings.append(f"{len(frags)} fragment(s) grammatical(aux) : " + " | ".join(frags[:3]))

    # 4. registre embarqué présent quand des lignes ont été retenues
    withheld = len((result or {}).get("v62136_withheld_scenarios") or [])
    metrics["withheld_lines"] = withheld
    has_register = bool(_RHR49.search(joined))                  # v62.13.9.4.49 — anciens et nouveaux titres
    metrics["register_present"] = has_register
    # v62.13.9.4.50 — en mode statut, rien n'est retenu hors du corps : les passages sont signalés en place et le
    # registre complet est le fichier probatoire (décision Robin du 30/09/2026 : plus d'annexe du registre dans le Word)
    try:
        from utils.v62136_fr_path_maintenance import _status_mode as _stm50
        _flag_mode50 = bool(_stm50(result or {}))
    except Exception:
        _flag_mode50 = False
    if withheld and not has_register and not _flag_mode50:
        blocking.append(f"{withheld} ligne(s) retenue(s) sans registre embarqué dans le fichier")
    if "hors fichier" in joined or "off-file" in joined:
        warnings.append("mention « détail conservé hors fichier » : le registre doit embarquer l'original")

    # 5. tables amputées : rangée dont toutes les cellules utiles sont des placeholders
    hollow = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if len(cells) >= 2 and all(_QA_PLACEHOLDER_RE.search(c) and len(c) < 70 for c in cells[1:]):
                hollow += 1
    metrics["hollow_table_rows"] = hollow
    if hollow:
        warnings.append(f"{hollow} rangée(s) de tableau entièrement neutralisée(s)")

    # 6. v62.13.9.4.18 — parité des rangées d'action (raw → final), point 8 :
    # une rangée d'action porte une échéance ; on compte celles des sorties
    # brutes et celles du fichier final.
    _date_re = re.compile(r"\b(?:\d{1,2}/\d{1,2}/20\d{2}|20\d{2}-\d{2}-\d{2}|\d{1,2}\s+(?:janvier|f[ée]vrier|mars|avril|mai|juin|juillet|ao[uû]t|septembre|octobre|novembre|d[ée]cembre)\s+20\d{2})\b", re.I)
    raw_rows = 0
    outs = (result or {}).get("outputs") or {}
    if isinstance(outs, Mapping):
        for v in outs.values():
            if isinstance(v, str):
                raw_rows += sum(1 for ln in v.splitlines() if ln.strip().startswith("|") and _date_re.search(ln)
                                and not re.match(r"^\|\s*-{2,}", ln.strip()))
    final_rows = 0
    for tbl in doc.tables:
        for row in tbl.rows:
            if _date_re.search(" | ".join(c.text for c in row.cells)):
                final_rows += 1
    metrics["action_rows_raw"] = raw_rows
    metrics["action_rows_final"] = final_rows
    if raw_rows and final_rows < 0.8 * raw_rows:
        warnings.append(f"parité des rangées d'action : {final_rows} dans le fichier pour {raw_rows} dans les sorties brutes")

    return {"version": "quality-audit-v2", "blocking": blocking, "warnings": warnings, "metrics": metrics}
