"""v62.13.9.4.21 — TYPOLOGIE DES CONTENUS DE L'OUTIL, À LA CONSTRUCTION.

Contre-audit ChatGPT du 05/09/2026, point 1 : le neutraliseur attaquait les
métadonnées du produit (version de build, scores, seuils de qualification,
compteurs de contrôle, intitulés techniques) parce que rien ne distinguait,
dans le fichier, un texte émis par l'outil d'une assertion d'agent — et
chaque cas nouveau appelait une regex de plus.

Désormais un bloc émis par l'outil est TYPÉ au moment où il est écrit : ses
paragraphes portent le style « CoDIR System ». Les juges (gate de surface,
oracle, neutraliseur de dernier recours, audit de qualité) sautent ces
paragraphes par leur style, sans lire leur texte. Un agent ne peut pas poser
ce style : il n'écrit que du texte.
"""
from __future__ import annotations

SYSTEM_STYLE = "CoDIR System"
SYSTEM_STYLE_ID = "CoDIRSystem"


def ensure_system_style(doc) -> None:
    """Crée le style de paragraphe « CoDIR System » (hérite de Normal) s'il manque."""
    try:
        from docx.enum.style import WD_STYLE_TYPE
        styles = doc.styles
        try:
            styles[SYSTEM_STYLE]
            return
        except KeyError:
            pass
        st = styles.add_style(SYSTEM_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        try:
            st.base_style = styles["Normal"]
        except Exception:
            pass
        st.hidden = False
        st.quick_style = False
    except Exception:
        return


def mark_paragraph_system(paragraph) -> None:
    try:
        paragraph.style = paragraph.part.document.styles[SYSTEM_STYLE]
    except Exception:
        try:
            ensure_system_style(paragraph.part.document)
            paragraph.style = paragraph.part.document.styles[SYSTEM_STYLE]
        except Exception:
            pass


def snapshot(doc) -> tuple[int, int]:
    return len(doc.paragraphs), len(doc.tables)


def mark_since(doc, snap: tuple[int, int]) -> int:
    """Type en « CoDIR System » tout ce qui a été ajouté depuis ``snap``."""
    ensure_system_style(doc)
    n0, t0 = snap
    marked = 0
    for p in doc.paragraphs[n0:]:
        mark_paragraph_system(p); marked += 1
    for tbl in doc.tables[t0:]:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    mark_paragraph_system(p); marked += 1
    return marked


def is_system_paragraph(paragraph) -> bool:
    try:
        return (paragraph.style is not None and str(paragraph.style.name or "") == SYSTEM_STYLE)
    except Exception:
        return False


def is_system_xml_block(block_xml: str) -> bool:
    """Pour l'oracle, qui lit l'XML brut : un <w:p> dont le pStyle est le style système."""
    return ('w:val="CoDIRSystem"' in block_xml) or ('w:val="CoDIR System"' in block_xml)
