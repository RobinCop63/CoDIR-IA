from __future__ import annotations

import io
import os
from dataclasses import dataclass
from pathlib import Path
from typing import List, Tuple, Optional

import pandas as pd

@dataclass
class AttachmentText:
    filename: str
    mimetype: str
    size_bytes: int
    text: str

def _estimate_tokens(text: str) -> int:
    # Rough heuristic: ~4 chars per token (English-ish). French is similar enough for budgeting.
    return max(1, int(len(text) / 4))

def _read_txt_bytes(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")

def extract_text_from_upload(filename: str, mimetype: str, data: bytes) -> str:
    ext = (Path(filename).suffix or "").lower()

    # Plain text / markdown
    if ext in [".txt", ".md"] or (mimetype.startswith("text/") if mimetype else False):
        return _read_txt_bytes(data)

    # CSV
    if ext == ".csv":
        df = pd.read_csv(io.BytesIO(data))
        return df.to_csv(index=False)

    # XLSX
    if ext in [".xlsx", ".xlsm"]:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        parts = []
        for ws in wb.worksheets[:5]:
            parts.append(f"## Feuille: {ws.title}")
            rows = []
            max_rows = min(ws.max_row or 0, 200)
            max_cols = min(ws.max_column or 0, 30)
            for r in range(1, max_rows + 1):
                row = []
                for c in range(1, max_cols + 1):
                    v = ws.cell(row=r, column=c).value
                    row.append("" if v is None else str(v))
                rows.append(row)
            # render as CSV-like text
            for row in rows:
                parts.append(",".join([x.replace("\n"," ").replace("\r"," ") for x in row]))
            parts.append("")
        return "\n".join(parts).strip()

    # DOCX
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

    # PPTX
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        out = []
        for i, slide in enumerate(prs.slides, start=1):
            out.append(f"## Slide {i}")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    out.append(shape.text.strip())
            out.append("")
        return "\n".join(out).strip()

    # PDF (best effort, no OCR)
    if ext == ".pdf":
        import pdfplumber
        txt_parts = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for i, page in enumerate(pdf.pages[:20], start=1):
                txt_parts.append(f"## Page {i}")
                text = page.extract_text() or ""
                if text.strip():
                    txt_parts.append(text)
                # try basic tables
                try:
                    tables = page.extract_tables() or []
                    for t in tables[:3]:
                        for row in t[:30]:
                            txt_parts.append("\t".join([("" if x is None else str(x)) for x in row]))
                except Exception:
                    pass
                txt_parts.append("")
        return "\n".join(txt_parts).strip()

    # Unknown/binary: return empty to avoid garbage
    return ""

def prepare_attachments(
    uploads: List[Tuple[str, str, bytes]],
    max_files: int = 5,
    max_total_mb: int = 25,
) -> List[AttachmentText]:
    uploads = uploads[:max_files]
    max_total_bytes = int(max_total_mb * 1024 * 1024)
    total = 0
    out: List[AttachmentText] = []
    for filename, mimetype, data in uploads:
        size = len(data)
        if total + size > max_total_bytes:
            break
        total += size
        text = extract_text_from_upload(filename, mimetype or "", data)
        out.append(AttachmentText(filename=filename, mimetype=mimetype or "", size_bytes=size, text=text))
    return out

def build_attachments_block(items: List[AttachmentText], token_budget: int = 8000) -> Tuple[str, dict]:
    """Return (text_block, stats). If too large, we truncate per file with head+tail."""
    if not items:
        return "", {"files": 0, "tokens_est": 0, "truncated": False}

    blocks = []
    total_tokens = 0
    truncated = False

    for a in items:
        t = a.text or ""
        tok = _estimate_tokens(t)
        # If adding full text exceeds budget, truncate this file.
        if total_tokens + tok > token_budget:
            truncated = True
            # keep head and tail
            keep_tokens = max(200, int(token_budget / max(1, len(items))))
            keep_chars = keep_tokens * 4
            head = t[:keep_chars]
            tail = t[-keep_chars:] if len(t) > keep_chars else ""
            t = (head + "\n\n[...]\n\n" + tail).strip()
            tok = _estimate_tokens(t)
        total_tokens += tok
        blocks.append(
            f"""### Fichier: {a.filename}\nType: {a.mimetype or 'n/a'} | Taille: {a.size_bytes} octets\n---\n{t}\n"""
        )

    # If still too large, hard cut from the end.
    joined = "\n\n".join(blocks).strip()
    joined_tokens = _estimate_tokens(joined)
    if joined_tokens > token_budget:
        truncated = True
        joined = joined[: token_budget * 4].rstrip() + "\n\n[...] (troncature globale)\n"
        joined_tokens = _estimate_tokens(joined)

    return joined, {"files": len(items), "tokens_est": joined_tokens, "truncated": truncated}
