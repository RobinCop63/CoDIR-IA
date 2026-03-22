"""
utils/pptx_export.py  —  CoDIR IA v16.3.3.22
=============================================
PowerPoint professionnel (palette Navy/Or) depuis le CR CoDIR IA.
API publique rétrocompatible avec app.py v16.3.3.22.
"""
from __future__ import annotations
import json, re
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ── Constantes métier ────────────────────────────────────────────────────────
REQUIRED_TOP_KEYS      = {"deck_title","date","topic","slides"}
MIN_TOPIC_SLIDES       = 4
MAX_TOPIC_SLIDES       = 7
MIN_BULLETS_PER_SLIDE  = 3
MAX_BULLETS_PER_SLIDE  = 5
ACTION_PLAN_MAX_ITEMS  = 6
TITLE_CONTEXT_FR = "Contexte & objectifs"
TITLE_ACTION_FR  = "Plan d'action (synthèse)"
TITLE_CONTEXT_EN = "Context & objectives"
TITLE_ACTION_EN  = "Action plan (summary)"
PREFERRED_DIRECTIONS_ORDER = ["strategie","marketing","finance","dg"]
DIRECTION_SYNONYMS = {
    "strategie":["strategie","stratégie","strat"],
    "marketing":["marketing"],
    "finance":  ["finance","financier","daf","cfo"],
    "dg":       ["dg","direction générale","direction generale","general management"],
}
_FORBIDDEN_PATTERNS = [
    r"tokens?_est",r"cache_(?:creation|read)_input_tokens",
    r"output\s*:\s*['\"]?\s*['\"]?$",r"source\s+pour\s+synth[eè]se\s+ppt",
    r"compte\s+rendu\s+codir\s+ia\s+—\s+source",r"pi[eè]ces\s+jointes",
    r"diagnostics?\s+ia",r"provider\s*:",r"model\s*:",r"elapsed",r"tronqu[eé]",
]

# ── Palette ──────────────────────────────────────────────────────────────────
C_DARK_NAVY  = RGBColor(0x14,0x1A,0x4A)
C_NAVY       = RGBColor(0x1E,0x27,0x61)
C_GOLD       = RGBColor(0xE8,0xB8,0x4B)
C_WHITE      = RGBColor(0xFF,0xFF,0xFF)
C_LIGHT_GRAY = RGBColor(0xF5,0xF7,0xFA)
C_ICE_BLUE   = RGBColor(0xCA,0xDC,0xFC)
C_MID_GRAY   = RGBColor(0x88,0x96,0xB0)
C_TEXT_DARK  = RGBColor(0x1A,0x1F,0x36)
C_GREEN      = RGBColor(0x27,0xAE,0x60)
C_TEAL       = RGBColor(0x0D,0x94,0x88)
C_RED        = RGBColor(0xE7,0x4C,0x3C)
C_BLUE_MED   = RGBColor(0x5E,0x81,0xF4)
C_ORANGE     = RGBColor(0xF9,0x61,0x67)
_DIR_COLORS  = {"strategie":C_TEAL,"marketing":C_ORANGE,"finance":C_GREEN,"dg":C_GOLD}
_DIR_LABELS  = {"strategie":"Stratégie","marketing":"Marketing","finance":"Finance","dg":"DG"}
TOPIC_COLORS = [C_TEAL,C_BLUE_MED,C_GREEN,C_ORANGE,C_RED,C_GOLD,C_NAVY]

# ── Helpers de rendu ─────────────────────────────────────────────────────────
def _add_rect(slide,x,y,w,h,fill,line=None):
    s=slide.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line: s.line.color.rgb=line
    else:    s.line.fill.background()
    return s

def _add_oval(slide,x,y,w,h,fill):
    s=slide.shapes.add_shape(9,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.fill.background(); return s

def _add_text(slide,text,x,y,w,h,size=12,bold=False,italic=False,
              color=None,align=PP_ALIGN.LEFT,font="Calibri"):
    if color is None: color=C_TEXT_DARK
    b=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=b.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.size=Pt(size)
    r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb=color; r.font.name=font

def _trunc(t,n): return t if len(t)<=n else t[:n-1]+"…"

def _header(slide,icon,title,subtitle):
    _add_rect(slide,0,0,10,1.08,C_NAVY)
    _add_rect(slide,0,1.08,10,0.055,C_GOLD)
    _add_rect(slide,0.22,0.2,0.65,0.65,C_GOLD)
    _add_text(slide,icon,0.22,0.2,0.65,0.65,size=22,bold=True,color=C_DARK_NAVY,align=PP_ALIGN.CENTER)
    _add_text(slide,title,1.05,0.18,8.6,0.48,size=22,bold=True,color=C_WHITE)
    _add_text(slide,subtitle,1.05,0.62,8.6,0.3,size=11,color=C_ICE_BLUE)

def _pgnum(slide,n):
    _add_text(slide,f"{n:02d}",9.5,5.3,0.45,0.25,size=10,color=C_MID_GRAY,align=PP_ALIGN.RIGHT)

# ── Slides spécialisées ──────────────────────────────────────────────────────
def _slide_title(prs,deck_title,topic,date_str):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=C_DARK_NAVY
    _add_rect(slide,0,0,0.1,5.625,C_GOLD)
    _add_oval(slide,6.8,-1.5,5.5,5.5,C_NAVY)
    _add_rect(slide,0.3,0.95,3.0,0.38,C_GOLD)
    _add_text(slide,"COMPTE-RENDU CoDIR IA",0.3,0.95,3.0,0.38,size=10,bold=True,color=C_DARK_NAVY,align=PP_ALIGN.CENTER)
    _add_text(slide,_trunc(deck_title,55),0.3,1.48,9.2,1.1,size=38,bold=True,color=C_WHITE,font="Calibri")
    _add_text(slide,_trunc(topic,75),0.3,2.65,9.0,0.55,size=20,bold=True,color=C_GOLD,font="Calibri")
    _add_rect(slide,0.3,3.35,5.5,0.04,C_MID_GRAY)
    _add_text(slide,"CoDIR IA — Comité de Direction augmenté par l'IA",0.3,3.52,8.0,0.38,size=13,bold=True,color=C_ICE_BLUE)
    _add_text(slide,"Stratégie · Marketing · Finance · Direction Générale",0.3,3.88,8.0,0.32,size=11,color=C_MID_GRAY)
    _add_rect(slide,0.3,4.72,2.2,0.5,C_NAVY)
    _add_text(slide,date_str,0.3,4.72,2.2,0.5,size=12,bold=True,color=C_ICE_BLUE,align=PP_ALIGN.CENTER)
    _add_text(slide,"v16.3.3.22",8.6,5.28,1.35,0.25,size=9,italic=True,color=C_MID_GRAY,align=PP_ALIGN.RIGHT)

def _slide_context(prs,title,bullets,page_n):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=C_LIGHT_GRAY
    _header(slide,"🎯",title,"Synthèse du brief initial · Recommandation")
    _add_rect(slide,0.28,1.25,9.44,4.1,C_WHITE)
    _add_rect(slide,0.28,1.25,0.07,4.1,C_GOLD)
    by=1.38
    for b in bullets[:MAX_BULLETS_PER_SLIDE]:
        lines=b.split("\n"); head=_trunc(lines[0],95); sub=_trunc(lines[1],95) if len(lines)>1 else ""
        _add_oval(slide,0.42,by+0.08,0.18,0.18,C_GOLD)
        _add_text(slide,head,0.72,by,8.8,0.32,size=13,bold=True,color=C_TEXT_DARK)
        if sub: _add_text(slide,sub,0.72,by+0.3,8.8,0.26,size=11,italic=True,color=C_MID_GRAY); by+=0.72
        else:   by+=0.72
    _pgnum(slide,page_n)

def _slide_topic(prs,title,bullets,idx,page_n):
    accent=TOPIC_COLORS[idx%len(TOPIC_COLORS)]
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=C_LIGHT_GRAY
    _header(slide,"📋",_trunc(title,60),f"Slide {page_n}")
    _add_rect(slide,0.28,1.25,9.44,4.1,C_WHITE)
    _add_rect(slide,0.28,1.25,0.08,4.1,accent)
    by=1.38
    for i,b in enumerate(bullets[:MAX_BULLETS_PER_SLIDE]):
        lines=b.split("\n"); head=_trunc(lines[0],95); sub=_trunc(lines[1],95) if len(lines)>1 else ""
        _add_rect(slide,0.42,by+0.02,0.28,0.28,accent)
        _add_text(slide,str(i+1),0.42,by+0.02,0.28,0.28,size=11,bold=True,color=C_WHITE,align=PP_ALIGN.CENTER)
        _add_text(slide,head,0.82,by,8.7,0.3,size=13,bold=True,color=C_TEXT_DARK)
        if sub: _add_text(slide,sub,0.82,by+0.3,8.7,0.26,size=11,italic=True,color=C_MID_GRAY); by+=0.75
        else:   by+=0.72
    _pgnum(slide,page_n)

def _slide_action(prs,title,bullets,page_n):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=C_DARK_NAVY
    _add_rect(slide,0,0,0.1,5.625,C_GOLD)
    _add_oval(slide,7.0,-1.2,5.5,5.5,C_NAVY)
    _add_rect(slide,0.3,0.15,2.8,0.42,C_GOLD)
    _add_text(slide,"PLAN D'ACTION",0.3,0.15,2.8,0.42,size=11,bold=True,color=C_DARK_NAVY,align=PP_ALIGN.CENTER)
    _add_text(slide,_trunc(title,55),0.3,0.68,9.0,0.48,size=24,bold=True,color=C_WHITE,font="Calibri")
    _add_text(slide,"1 décision par direction — Ordre : Stratégie · Marketing · Finance · DG",
              0.3,1.22,9.0,0.3,size=11,italic=True,color=C_ICE_BLUE)
    _add_rect(slide,0.3,1.6,9.4,0.04,C_NAVY)
    by=1.75
    for b in bullets[:ACTION_PLAN_MAX_ITEMS]:
        dk=_direction_key_from_bullet(b); dc=_DIR_COLORS.get(dk,C_GOLD); dl=_DIR_LABELS.get(dk,"—")
        _add_rect(slide,0.3,by,9.4,0.52,C_NAVY)
        _add_rect(slide,0.3,by,0.06,0.52,dc)
        _add_rect(slide,0.42,by+0.06,1.15,0.36,dc)
        _add_text(slide,dl,0.42,by+0.06,1.15,0.36,size=10,bold=True,color=C_DARK_NAVY,align=PP_ALIGN.CENTER)
        _add_text(slide,_trunc(b,100),1.68,by+0.1,7.9,0.32,size=12,color=C_WHITE)
        by+=0.6
    _pgnum(slide,page_n)

# ── Point d'entrée ───────────────────────────────────────────────────────────
def create_pptx_from_deck(payload: Dict[str,Any], out_path: str) -> str:
    ok,msg=validate_deck(payload)
    if not ok: raise ValueError(msg)
    prs=Presentation(); prs.slide_width=Inches(10); prs.slide_height=Inches(5.625)
    _slide_title(prs,str(payload.get("deck_title","CoDIR IA — Synthèse")),
                 str(payload.get("topic","")),str(payload.get("date",date.today().isoformat())))
    slides=payload["slides"]; tidx=0
    for i,sld in enumerate(slides):
        title=str(sld.get("title","")).strip()
        bullets=[str(b) for b in (sld.get("bullets") or []) if str(b).strip()]
        pn=i+2
        if i==0:               _slide_context(prs,title,bullets,pn)
        elif i==len(slides)-1: _slide_action(prs,title,bullets,pn)
        else:                  _slide_topic(prs,title,bullets,tidx,pn); tidx+=1
    prs.save(out_path); return out_path

# ── Nettoyage texte ──────────────────────────────────────────────────────────
_FR_STOP={"le","la","les","de","des","du","un","une","et","ou","pour","avec","sans","sur","dans","ce","cette","ces","au","aux","en","que","qui","quoi","dont","plus","moins","très","afin","mais","donc","car"}
_EN_STOP={"the","a","an","and","or","to","for","with","without","on","in","this","that","these","those","as","at","by","from","of","is","are","be","will","should","can","may","more","less"}

def _detect_language(text):
    t=re.sub(r"[^A-Za-zÀ-ÿ\s]"," ",(text or "")).lower()
    words=[w for w in t.split() if len(w)>=2][:500]
    if not words: return "fr"
    fr=sum(1 for w in words if w in _FR_STOP); en=sum(1 for w in words if w in _EN_STOP)
    return "en" if en>=fr*1.4 and en>=6 else "fr"

def _is_markdown_table_line(t): return t.count("|")>=2 and not t.strip().startswith("http")

def _clean_line(t):
    t=(t or "").strip()
    if not t or t in {"---","—","–"} or t.startswith("#"): return ""
    if _is_markdown_table_line(t): return ""
    t=re.sub(r"`+","",t).strip()
    if t.startswith("{") and ("provider" in t or "model" in t or "tokens" in t): return ""
    low=t.lower()
    for pat in _FORBIDDEN_PATTERNS:
        if re.search(pat,low,flags=re.IGNORECASE): return ""
    return t

def clean_report_for_ppt(report_text,max_chars=25000):
    lines,total=[],0
    for ln in (report_text or "").splitlines():
        t=_clean_line(ln)
        if not t: continue
        lines.append(t); total+=len(t)+1
        if total>max_chars: break
    return "\n".join(lines).strip()[:max_chars]

def extract_key_figures(clean_text,max_items=18):
    out,seen=[],set()
    for ln in (clean_text or "").splitlines():
        t=(ln or "").strip()
        if not t or not any(c.isdigit() for c in t): continue
        if any(re.search(p,t.lower(),flags=re.IGNORECASE) for p in _FORBIDDEN_PATTERNS): continue
        t=re.sub(r"\s+"," ",t).strip()
        if len(t)>180:
            m=re.search(r"\d",t)
            if m: t=t[max(0,m.start()-60):min(len(t),m.start()+120)].strip()
            t=t[:180].strip()
        if t and t not in seen: out.append(t); seen.add(t)
        if len(out)>=max_items: break
    return out

_BULLET_PREFIX_RE=re.compile(r"^\s*(?:[-•–*]+|\d+[\)\.]|[A-Za-z][\)\.]|[✅⚠️👉➡️⭐️🔹🔸]+)\s*")
_MD_BOLD_RE      =re.compile(r"(\*\*|__)(.*?)\1")
_MD_ITALIC_RE    =re.compile(r"(?<!\*)\*(?!\*)([^*]{2,}?)\*(?!\*)")
_MD_ITALIC2_RE   =re.compile(r"(?<!_)_(?!_)([^_]{2,}?)_(?!_)")

def _explode_listblock(b):
    raw=[ln.rstrip() for ln in (b or "").splitlines() if ln.strip()]
    if len(raw)<3: return [b]
    looks=sum(1 for ln in raw if _BULLET_PREFIX_RE.match(ln.strip()))
    return raw if looks>=2 else [b]

def _strip_md(t):
    t=re.sub(r"\[([^\]]+)\]\(([^\)]+)\)",r"\1",t)
    return re.sub(r"`+","",t).replace("**","").replace("__","").replace("*","")

def _clean_bullet_text(b,*,single_line=False):
    b=(b or "").strip()
    if not b: return ""
    b=_strip_md(b).strip()
    b=_MD_BOLD_RE.sub(r"\2",b); b=_MD_ITALIC_RE.sub(r"\1",b); b=_MD_ITALIC2_RE.sub(r"\1",b)
    raw=[ln.strip() for ln in b.splitlines() if ln.strip()]
    if not raw: return ""
    raw=raw[:1 if single_line else 2]; cleaned=[]
    for ln in raw:
        ln=re.sub(r"^\s*[-•–*]+\s*","",ln).strip()
        ln=_BULLET_PREFIX_RE.sub("",ln).strip()
        ln=re.sub(r"^[^0-9A-Za-zÀ-ÿ]+\s*","",ln).strip()
        if len(ln)>20 and ln.endswith("."): ln=ln[:-1].rstrip()
        ln=ln.replace("**","").replace("__","")
        ln=re.sub(r"\s+"," ",ln).strip()[:240]
        if ln: cleaned.append(ln)
    return "\n".join(cleaned).strip()

def _sentences(text):
    t=re.sub(r"\s+"," ",(text or "")).strip()
    if not t: return []
    out=[]
    for p in re.split(r"(?<=[\.!\?;])\s+",t):
        p=p.strip()
        if len(p)>=18 and p not in out: out.append(p)
    return out

def _extract_block(t,starts,ends):
    if not t: return ""
    low=t.lower()
    si=[i for sm in starts for i in [low.find(sm.lower())] if i!=-1]
    if not si: return ""
    s=min(si); after=low[s:]
    ei=[s+j for em in ends for j in [after.find(em.lower())] if j!=-1]
    return t[s:min(ei) if ei else len(t)].strip()

def _extract_initial_prompt(clean_text):
    for st,en in [
        (["BRIEF INITIAL","BRIEF","Brief initial","Brief"],["STRATÉGIE","STRATEGY","Constats"]),
        (["Pouvez vous","Peux-tu","Can you"],["STRATÉGIE","STRATEGY","Constats"]),
    ]:
        c=re.sub(r"\s+"," ",_extract_block(clean_text,st,en) or "").strip()
        if len(c)>=40: return c[:1200]
    return ""

def infer_topic_slide_count(clean_text):
    n=len(clean_text or ""); nl=len([ln for ln in (clean_text or "").splitlines() if ln.strip()])
    if n<4500 and nl<90:   return 4
    if n<9000 and nl<180:  return 5
    if n<14000 and nl<280: return 6
    return 7

def _direction_key_from_bullet(bullet):
    b=re.sub(r"[^a-z0-9à-ÿ\s]"," ",(bullet or "").lower()); b=re.sub(r"\s+"," ",b).strip()
    for key in PREFERRED_DIRECTIONS_ORDER:
        for kw in DIRECTION_SYNONYMS.get(key,[]):
            if kw in b: return key
    return ""

# ── Validation ───────────────────────────────────────────────────────────────
def validate_deck(payload):
    if not isinstance(payload,dict): return False,"Le JSON doit être un objet."
    missing=REQUIRED_TOP_KEYS-set(payload.keys())
    if missing: return False,f"Clés manquantes: {', '.join(sorted(missing))}"
    slides=payload.get("slides")
    if not isinstance(slides,list): return False,"slides doit être une liste."
    if len(slides)<1+MIN_TOPIC_SLIDES+1 or len(slides)>1+MAX_TOPIC_SLIDES+1:
        return False,"Nombre de slides invalide."
    lang=_detect_language(" ".join([str(payload.get("deck_title","")),str(payload.get("topic",""))]))
    ctx=TITLE_CONTEXT_EN if lang=="en" else TITLE_CONTEXT_FR
    act=TITLE_ACTION_EN  if lang=="en" else TITLE_ACTION_FR
    if (slides[0].get("title") or "").strip().lower()!=ctx.lower():
        return False,f"La première slide doit s'intituler '{ctx}'."
    if (slides[-1].get("title") or "").strip().lower()!=act.lower():
        return False,f"La dernière slide doit s'intituler '{act}'."
    for i,sld in enumerate(slides,start=1):
        if not isinstance(sld,dict): return False,f"Slide {i}: objet requis."
        if "title" not in sld or "bullets" not in sld: return False,f"Slide {i}: title+bullets requis."
        if not (sld.get("title") or "").strip(): return False,f"Slide {i}: title vide."
        buls=sld.get("bullets")
        if not isinstance(buls,list) or not all(isinstance(b,str) for b in buls):
            return False,f"Slide {i}: bullets=[str]."
        is_act=(i==len(slides)); min_b=1 if is_act else MIN_BULLETS_PER_SLIDE; max_b=ACTION_PLAN_MAX_ITEMS if is_act else MAX_BULLETS_PER_SLIDE
        cleaned=[]
        for b in buls:
            for b_raw in _explode_listblock(b):
                b2=_clean_bullet_text(b_raw,single_line=is_act)
                if not b2: continue
                if _is_markdown_table_line(b2): return False,f"Slide {i}: table Markdown."
                if any(re.search(p,b2,flags=re.IGNORECASE) for p in _FORBIDDEN_PATTERNS): return False,f"Slide {i}: bruit technique."
                cleaned.append(b2)
        if len(cleaned)<min_b: return False,f"Slide {i}: >= {min_b} bullets requis."
        if len(cleaned)>max_b: return False,f"Slide {i}: <= {max_b} bullets."
    return True,"OK"

def coerce_deck(payload):
    if not isinstance(payload,dict): return payload
    slides=payload.get("slides") if isinstance(payload.get("slides"),list) else []
    lang=_detect_language(" ".join([str(payload.get("deck_title","")),str(payload.get("topic","")),json.dumps(slides,ensure_ascii=False)[:2000]]))
    ctx=TITLE_CONTEXT_EN if lang=="en" else TITLE_CONTEXT_FR
    act=TITLE_ACTION_EN  if lang=="en" else TITLE_ACTION_FR
    desired=max(MIN_TOPIC_SLIDES,min(MAX_TOPIC_SLIDES,infer_topic_slide_count(json.dumps(slides,ensure_ascii=False))))

    def _cb(raw,single_line=False):
        out=[]
        for b in raw:
            if not isinstance(b,str): continue
            for b_raw in _explode_listblock(b):
                b2=_clean_bullet_text(b_raw,single_line=single_line)
                if b2 and not _is_markdown_table_line(b2) and not any(re.search(p,b2,flags=re.IGNORECASE) for p in _FORBIDDEN_PATTERNS):
                    out.append(b2)
        return out

    s0=slides[0] if slides and isinstance(slides[0],dict) else {}
    b0=_cb(s0.get("bullets") or [])[:MAX_BULLETS_PER_SLIDE]
    kw=["recommandation","critère de succès","recommendation","success criteria"]
    if not any(any(k in (bb or "").lower() for k in kw) for bb in b0):
        reco="Recommendation / success criteria: To be confirmed" if lang=="en" else "Recommandation / critère de succès\nÀ confirmer"
        b0=(b0+[reco])[:MAX_BULLETS_PER_SLIDE]
    fixed=[{"title":ctx,"bullets":b0}]

    src=slides[1:-1] if len(slides)>=3 else slides[1:]
    for i in range(desired):
        sld=src[i] if i<len(src) and isinstance(src[i],dict) else {}
        title=re.sub(r"[\n\r\t]+"," ",(sld.get("title") or "").strip())[:80] or ("Topic" if lang=="en" else "Sujet")
        bl=_cb(sld.get("bullets") or [])[:MAX_BULLETS_PER_SLIDE]
        fixed.append({"title":title,"bullets":bl})

    slast=slides[-1] if len(slides)>1 and isinstance(slides[-1],dict) else {}
    b_last=_cb(slast.get("bullets") or [],single_line=True)
    def _sk(b): k=_direction_key_from_bullet(b); return (0,PREFERRED_DIRECTIONS_ORDER.index(k)) if k in PREFERRED_DIRECTIONS_ORDER else (1,b.lower())
    b_last=sorted(dict.fromkeys(b_last),key=_sk)[:ACTION_PLAN_MAX_ITEMS]
    fixed.append({"title":act,"bullets":b_last})

    for i,sld in enumerate(fixed):
        is_act=(i==len(fixed)-1); min_b=1 if is_act else MIN_BULLETS_PER_SLIDE
        filler="À compléter" if lang!="en" else "To be completed"
        while len(sld["bullets"])<min_b: sld["bullets"].append(filler)

    payload["slides"]     =fixed[:1+desired+1]
    payload["deck_title"] =(payload.get("deck_title") or ("CoDIR IA — Summary" if lang=="en" else "CoDIR IA — Synthèse")).strip()
    payload["topic"]      =(payload.get("topic") or ("CoDIR meeting report summary" if lang=="en" else "Synthèse du compte rendu CoDIR")).strip()
    payload["date"]       =(payload.get("date") or date.today().isoformat()).strip()
    return payload

def fallback_deck_from_report(report_text):
    clean=clean_report_for_ppt(report_text or "",max_chars=18000)
    lang=_detect_language(clean)
    ctx=TITLE_CONTEXT_EN if lang=="en" else TITLE_CONTEXT_FR
    act=TITLE_ACTION_EN  if lang=="en" else TITLE_ACTION_FR
    topic_n=max(MIN_TOPIC_SLIDES,min(MAX_TOPIC_SLIDES,infer_topic_slide_count(clean)))
    sents=_sentences(clean); ctx_b=[]
    for s in sents[:6]:
        for r in _explode_listblock(s):
            b=_clean_bullet_text(r)
            if b: ctx_b.append(b)
        if len(ctx_b)>=MAX_BULLETS_PER_SLIDE: break
    if not any(("recommandation" in (b or "").lower() or "recommendation" in (b or "").lower()) for b in ctx_b):
        ctx_b.append("Recommandation / critère de succès\nÀ confirmer" if lang!="en" else "Recommendation / success criteria\nTo be confirmed")
    ctx_b=ctx_b[:MAX_BULLETS_PER_SLIDE]
    lines=[ln.strip() for ln in clean.splitlines() if ln.strip()]
    chunks,ct,cl=[],None,[]
    for ln in lines:
        ih=(ln.endswith(":") and len(ln)<80) or (ln.isupper() and 6<=len(ln)<=60)
        if ih and cl: chunks.append((ct or ("Topic" if lang=="en" else "Sujet")," ".join(cl))); ct=ln.rstrip(":").strip(); cl=[]
        elif ih: ct=ln.rstrip(":").strip()
        else: cl.append(ln)
    if cl: chunks.append((ct or ("Topic" if lang=="en" else "Sujet")," ".join(cl)))
    ts=[]
    for i in range(topic_n):
        tt,tb=chunks[i] if i<len(chunks) else (("Topic" if lang=="en" else "Sujet"),clean)
        tt=re.sub(r"[\n\r\t]+"," ",(tt or "")).strip()[:80]
        bl=[]
        for s in _sentences(tb):
            b=_clean_bullet_text(s)
            if b: bl.append(b)
            if len(bl)>=MAX_BULLETS_PER_SLIDE: break
        while len(bl)<MIN_BULLETS_PER_SLIDE: bl.append("À compléter" if lang!="en" else "To be completed")
        ts.append({"title":tt,"bullets":bl[:MAX_BULLETS_PER_SLIDE]})
    act_c=[ln for ln in lines if any(k in ln.lower() for k in ["plan d'action","actions","next","roadmap","to do"])]
    if not act_c: act_c=extract_key_figures(clean,max_items=ACTION_PLAN_MAX_ITEMS) or lines[:ACTION_PLAN_MAX_ITEMS]
    ab=[]
    for a in act_c:
        for r in _explode_listblock(a):
            b=_clean_bullet_text(r,single_line=True)
            if b: ab.append(b)
            if len(ab)>=ACTION_PLAN_MAX_ITEMS: break
    if not ab: ab=["À compléter" if lang!="en" else "To be completed"]
    def _sk(b): k=_direction_key_from_bullet(b); return (0,PREFERRED_DIRECTIONS_ORDER.index(k)) if k in PREFERRED_DIRECTIONS_ORDER else (1,b.lower())
    ab=sorted(dict.fromkeys(ab),key=_sk)[:ACTION_PLAN_MAX_ITEMS]
    return {"deck_title":"CoDIR IA — Summary" if lang=="en" else "CoDIR IA — Synthèse",
            "topic":"CoDIR meeting report summary" if lang=="en" else "Synthèse du compte rendu CoDIR",
            "date":date.today().isoformat(),
            "slides":[{"title":ctx,"bullets":ctx_b}]+ts+[{"title":act,"bullets":ab}]}

# ── Prompt IA & JSON ─────────────────────────────────────────────────────────
def build_pptx_prompt(report_text):
    clean=clean_report_for_ppt(report_text,max_chars=25000)
    lang=_detect_language(clean)
    topic_n=max(MIN_TOPIC_SLIDES,min(MAX_TOPIC_SLIDES,infer_topic_slide_count(clean)))
    total=1+topic_n+1
    ctx=TITLE_CONTEXT_EN if lang=="en" else TITLE_CONTEXT_FR
    act=TITLE_ACTION_EN  if lang=="en" else TITLE_ACTION_FR
    ip=_extract_initial_prompt(clean); kf=extract_key_figures(clean)
    if lang=="en":
        hint=(f"You generate STRICT valid JSON for a PowerPoint summary.\nReturn ONLY JSON.\n"
              f"- slides MUST have exactly {total} items.\n"
              f"- slides[0].title MUST be exactly '{ctx}'.\n"
              f"- slides[-1].title MUST be exactly '{act}'.\n"
              f"- Topic slides: {MIN_BULLETS_PER_SLIDE}-{MAX_BULLETS_PER_SLIDE} bullets each.\n"
              "- Bullets: NO prefix, action verb (infinitive), 1 line + optional sub-line ('\\n').\n"
              "- Keep key figures (%, EUR, dates, KPIs). Do NOT invent numbers.\n"
              f"- Action plan: 1 bullet per function (max {ACTION_PLAN_MAX_ITEMS}), single-line. Order: Strategy, Marketing, Finance, DG.\n"
              'JSON schema: {"deck_title":str,"topic":str,"date":str,"slides":[{"title":str,"bullets":[str]}]}\n')
    else:
        hint=(f"Tu génères un JSON STRICTEMENT valide pour un PowerPoint.\nRéponds UNIQUEMENT en JSON.\n"
              f"- slides DOIT contenir exactement {total} éléments.\n"
              f"- slides[0].title DOIT être exactement : '{ctx}'.\n"
              f"- slides[-1].title DOIT être exactement : '{act}'.\n"
              f"- Slides sujets : {MIN_BULLETS_PER_SLIDE}-{MAX_BULLETS_PER_SLIDE} bullets.\n"
              "- Chaque bullet : SANS préfixe, verbe d'action infinitif, 1 ligne + 1 optionnelle ('\\n').\n"
              "- Conserver chiffres clés (%, €, dates, KPI). Ne jamais inventer.\n"
              f"- Plan d'action : 1 bullet par direction (max {ACTION_PLAN_MAX_ITEMS}), 1 seule ligne. Ordre : Stratégie, Marketing, Finance, DG.\n"
              'Schéma JSON : {"deck_title":str,"topic":str,"date":str,"slides":[{"title":str,"bullets":[str]}]}\n')
    body=(f"\n---\n{'CONTENT' if lang=='en' else 'CONTENU'}\n"+clean[:25000]
          +("\n\n---\n"+("KEY FIGURES" if lang=="en" else "CHIFFRES CLÉS")+"\n"+"\n".join(kf) if kf else "")
          +("\n\n---\n"+("INITIAL PROMPT" if lang=="en" else "PROMPT INITIAL")+"\n"+ip if ip else ""))
    return hint+body

def _extract_json_object(text):
    if not text: return ""
    s=text.find("{"); e=text.rfind("}")
    if s==-1 or e==-1 or e<=s: return ""
    return text[s:e+1].strip()

def parse_json_relaxed(text):
    js=_extract_json_object((text or "").strip())
    if not js: raise ValueError("JSON introuvable.")
    js=re.sub(r",\s*([}\]])",r"\1",js)
    return json.loads(js)

def parse_json_strict(text):
    t=(text or "").strip()
    if not t.startswith("{"): t=_extract_json_object(t)
    if not t: raise ValueError("JSON introuvable.")
    return json.loads(t)
